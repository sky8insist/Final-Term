from __future__ import annotations

import json
from pathlib import Path

import fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw

from experiments.mineru_ab.config import EXPERIMENT_ROOT


GENERATED = EXPERIMENT_ROOT / "fixtures" / "generated"
MANIFEST = EXPERIMENT_ROOT / "fixtures" / "manifest.json"


def _make_pdf(path: Path, pages: list[list[tuple[str, float]]]) -> None:
    document = fitz.open()
    for lines in pages:
        page = document.new_page()
        y = 72.0
        for text, size in lines:
            page.insert_text((72, y), text, fontsize=size, fontname="helv")
            y += size + 14
    document.save(path)
    document.close()


def _make_docx(path: Path) -> None:
    document = Document()
    if "Experiment Heading" not in [style.name for style in document.styles]:
        document.styles.add_style("Experiment Heading", WD_STYLE_TYPE.PARAGRAPH)
    document.add_heading("Network Review", level=1)
    document.add_paragraph("TCP slow start increases the congestion window during early transmission.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Protocol"
    table.cell(0, 1).text = "Layer"
    table.cell(1, 0).text = "TCP"
    table.cell(1, 1).text = "Transport"
    document.save(path)


def _make_scanned_pdf(path: Path) -> None:
    image = Image.new("RGB", (1200, 500), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 180), "SCANNED IMAGE WITHOUT A PDF TEXT LAYER", fill="black")
    image_path = path.with_suffix(".png")
    image.save(image_path)
    document = fitz.open()
    page = document.new_page(width=600, height=250)
    page.insert_image(page.rect, filename=str(image_path))
    document.save(path)
    document.close()
    image_path.unlink()


def generate() -> list[dict]:
    GENERATED.mkdir(parents=True, exist_ok=True)
    _make_pdf(GENERATED / "simple_text.pdf", [[
        ("Computer Networks", 20),
        ("TCP provides reliable ordered byte-stream delivery.", 12),
        ("Slow start grows the congestion window during early transmission.", 12),
    ]])
    _make_pdf(GENERATED / "two_pages.pdf", [
        [("Chapter One", 20), ("A router forwards packets between networks.", 12)],
        [("Chapter Two", 20), ("HTTP is an application-layer protocol.", 12)],
    ])
    _make_docx(GENERATED / "table_document.docx")
    _make_scanned_pdf(GENERATED / "scanned_no_text.pdf")
    cases = [
        {
            "id": "simple_text_pdf", "path": "generated/simple_text.pdf",
            "content_type": "application/pdf", "expected_outcome": "success",
            "expected_text": "Computer Networks TCP provides reliable ordered byte-stream delivery. Slow start grows the congestion window during early transmission.",
        },
        {
            "id": "two_page_pdf", "path": "generated/two_pages.pdf",
            "content_type": "application/pdf", "expected_outcome": "success",
            "expected_text": "Chapter One A router forwards packets between networks. Chapter Two HTTP is an application-layer protocol.",
        },
        {
            "id": "table_docx", "path": "generated/table_document.docx",
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "expected_outcome": "success",
            "expected_text": "Network Review TCP slow start increases the congestion window during early transmission. Protocol Layer TCP Transport",
        },
        {
            "id": "scanned_pdf", "path": "generated/scanned_no_text.pdf",
            "content_type": "application/pdf", "expected_outcome": "ocr_required",
            "expected_text": "",
        },
    ]
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps({"cases": cases}, ensure_ascii=False, indent=2), encoding="utf-8")
    return cases


if __name__ == "__main__":
    generated = generate()
    print(f"Generated {len(generated)} isolated experiment cases in {GENERATED}")

