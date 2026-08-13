from io import BytesIO

import pytest
from docx import Document

from app.services.chunk_service import build_block_chunk_records, build_chunk_records, chunk_text
from app.services.parse_service import (
    DOCX_CONTENT_TYPE,
    DocumentParseError,
    PDF_CONTENT_TYPE,
    TXT_CONTENT_TYPE,
    parse_document_blocks,
    parse_document_bytes,
)


def test_parse_txt_supports_utf8_bom():
    text = parse_document_bytes(
        "\ufeff第一章\nmodule four".encode("utf-8"),
        filename="notes.txt",
        content_type=TXT_CONTENT_TYPE,
    )

    assert text == "第一章\nmodule four"


def test_parse_txt_detects_common_chinese_encoding():
    text = parse_document_bytes(
        "期末复习：矩阵与特征值".encode("gb18030"),
        filename="notes.txt",
        content_type=TXT_CONTENT_TYPE,
    )

    assert text == "期末复习：矩阵与特征值"


def test_parse_txt_as_versioned_content_block():
    blocks = parse_document_blocks(
        b"hello\nworld", filename="notes.txt", content_type=TXT_CONTENT_TYPE,
    )

    assert len(blocks) == 1
    assert blocks[0]["block_type"] == "paragraph"
    assert blocks[0]["content_text"] == "hello\nworld"
    assert blocks[0]["parser_name"] == "text"
    assert len(blocks[0]["source_hash"]) == 64


def test_parse_docx_extracts_paragraph_text():
    document = Document()
    document.add_paragraph("第一段")
    document.add_paragraph("second paragraph")
    buffer = BytesIO()
    document.save(buffer)

    text = parse_document_bytes(
        buffer.getvalue(),
        filename="notes.docx",
        content_type=DOCX_CONTENT_TYPE,
    )

    assert "第一段" in text
    assert "second paragraph" in text


def test_parse_pdf_uses_pypdf_extract_text(monkeypatch):
    class FakePage:
        def extract_text(self):
            return "pdf text"

    class FakeReader:
        def __init__(self, _stream):
            self.pages = [FakePage()]

    monkeypatch.setattr("app.services.parse_service.PdfReader", FakeReader)

    text = parse_document_bytes(
        b"%PDF-1.4",
        filename="notes.pdf",
        content_type=PDF_CONTENT_TYPE,
    )

    assert text == "pdf text"


def test_pdf_content_blocks_preserve_page_and_coordinates():
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Matrix eigenvalue review")
    data = document.tobytes()
    document.close()

    blocks = parse_document_blocks(data, filename="notes.pdf", content_type=PDF_CONTENT_TYPE)

    assert blocks[0]["page_number"] == 1
    assert blocks[0]["bounding_box"]["x0"] >= 0
    assert blocks[0]["bounding_box"]["pageWidth"] > 0


def test_parse_empty_pdf_fails(monkeypatch):
    class FakePage:
        def extract_text(self):
            return ""

    class FakeReader:
        def __init__(self, _stream):
            self.pages = [FakePage()]

    monkeypatch.setattr("app.services.parse_service.PdfReader", FakeReader)

    with pytest.raises(DocumentParseError, match="No extractable text"):
        parse_document_bytes(
            b"%PDF-1.4",
            filename="notes.pdf",
            content_type=PDF_CONTENT_TYPE,
        )


def test_chunk_overlap_is_stable():
    chunks = chunk_text("abcdefghij", chunk_size=4, overlap=2)

    assert chunks == ["abcd", "cdef", "efgh", "ghij"]


def test_build_chunk_records_adds_metadata():
    records = build_chunk_records("abcdefghij", chunk_size=4, overlap=2)

    assert records[0]["chunk_index"] == 0
    assert records[0]["content"] == "abcd"
    assert records[0]["metadata"]["chunk_size"] == 4
    assert records[0]["metadata"]["overlap"] == 2


def test_block_chunks_preserve_source_location_and_never_cross_blocks():
    records = build_block_chunk_records([
        {
            "id": "block-1", "block_type": "paragraph", "content_text": "abcdefgh",
            "page_number": 2, "bounding_box": {"x0": 1}, "start_time": None,
            "end_time": None, "sequence_index": 0, "parser_name": "pymupdf",
            "parser_version": "1.0.0", "confidence": 0.95,
        },
        {
            "id": "block-2", "block_type": "audio", "content_text": "ijkl",
            "page_number": None, "bounding_box": None, "start_time": 10.0,
            "end_time": 12.5, "sequence_index": 1, "parser_name": "transcription-api",
            "parser_version": "1.0.0", "confidence": 0.8,
        },
    ], chunk_size=5, overlap=1)

    assert [record["content_block_id"] for record in records] == ["block-1", "block-1", "block-2"]
    assert records[0]["page_number"] == 2
    assert records[0]["bounding_box"] == {"x0": 1}
    assert records[2]["start_time"] == 10.0
    assert records[2]["end_time"] == 12.5
    assert all("ijkl" not in record["content"] for record in records[:2])
