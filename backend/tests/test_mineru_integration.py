from io import BytesIO
import json
from zipfile import ZipFile

import httpx
from openpyxl import Workbook
from pptx import Presentation
from starlette.datastructures import Headers, UploadFile

from app.config.settings import settings
from app.services.mineru_adapter import read_mineru_bundle
from app.services.mineru_client import MinerUClient
from app.services import mineru_parser
from app.services.mineru_parser import _apply_native_quality_gate, parse_material, uses_mineru
from app.services.file_service import normalized_upload_content_type, validate_upload_file
from app.services.parse_service import parse_native_document_blocks


def _response(payload: dict) -> httpx.Response:
    return httpx.Response(200, json=payload)


def test_mineru_disables_ocr_for_documents_and_enables_it_for_images(monkeypatch):
    payloads = []

    def handler(request: httpx.Request) -> httpx.Response:
        payloads.append(json.loads(request.content))
        return _response({
            "code": 0, "trace_id": "trace",
            "data": {"batch_id": f"batch-{len(payloads)}", "file_urls": ["https://upload.example/file"]},
        })

    monkeypatch.setattr(settings, "mineru_api_token", "test-token")
    transport = httpx.MockTransport(handler)
    client = httpx.Client(transport=transport)
    provider = MinerUClient(client=client, transfer_client=client, sleep=lambda _: None)

    provider.request_upload(filename="lecture.pdf", data_id="pdf-1", is_ocr=False)
    provider.request_upload(filename="scan.png", data_id="image-1", is_ocr=True)

    assert payloads[0]["model_version"] == "pipeline"
    assert payloads[0]["files"][0]["is_ocr"] is False
    assert payloads[1]["files"][0]["is_ocr"] is True


def test_mineru_retries_transient_api_response(monkeypatch):
    calls = 0

    def handler(_request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        if calls == 1:
            return httpx.Response(503, text="temporarily unavailable")
        return _response({
            "code": 0,
            "data": {"batch_id": "batch-retry", "file_urls": ["https://upload.example/file"]},
        })

    monkeypatch.setattr(settings, "mineru_api_token", "test-token")
    monkeypatch.setattr(settings, "mineru_transport_retries", 2)
    client = httpx.Client(transport=httpx.MockTransport(handler))
    provider = MinerUClient(client=client, transfer_client=client, sleep=lambda _: None)

    ticket = provider.request_upload(filename="lecture.pdf", data_id="retry", is_ocr=False)

    assert calls == 2
    assert ticket.batch_id == "batch-retry"


def test_mineru_bundle_maps_stable_content_list_to_content_blocks():
    content = [
        {"type": "text", "text": "Chapter 1", "text_level": 1, "page_idx": 0, "bbox": [1, 2, 3, 4]},
        {"type": "table", "table_body": "| A | B |", "page_idx": 0},
    ]
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("demo_content_list.json", json.dumps(content))
        archive.writestr("demo_full.md", "# Chapter 1")

    result = read_mineru_bundle(buffer.getvalue())

    assert [block["block_type"] for block in result.blocks] == ["heading", "table"]
    assert result.blocks[0]["page_number"] == 1
    assert result.blocks[0]["bounding_box"] == {"x0": 1, "y0": 2, "x1": 3, "y1": 4}
    assert result.markdown == b"# Chapter 1"


def test_resume_batch_does_not_request_or_upload_file(monkeypatch):
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("demo_content_list.json", json.dumps([
            {"type": "text", "text": "resumed result", "page_idx": 0},
        ]))

    class FakeClient:
        def __enter__(self): return self
        def __exit__(self, *_args): return None
        def request_upload(self, **_kwargs): raise AssertionError("must not resubmit")
        def upload_bytes(self, *_args): raise AssertionError("must not reupload")
        def wait_for_result(self, batch_id, on_state=None):
            assert batch_id == "existing-batch"
            if on_state: on_state("done")
            return {"full_zip_url": "https://result.example/result.zip"}
        def download_result(self, _url): return buffer.getvalue()

    monkeypatch.setattr(mineru_parser, "MinerUClient", FakeClient)
    monkeypatch.setattr(settings, "enable_mineru", True)

    result = parse_material(
        b"legacy-office", filename="legacy.doc", content_type="application/msword",
        material_id="material-1", resume_batch_id="existing-batch",
    )

    assert result.provider == "mineru-v4"
    assert result.metadata["batchId"] == "existing-batch"


def test_quality_gate_replaces_only_incomplete_page_text(monkeypatch):
    monkeypatch.setattr(settings, "mineru_page_coverage_threshold", 0.9)
    primary = [
        {"block_type": "paragraph", "content_text": "short", "page_number": 1, "sequence_index": 0, "metadata": {}},
        {"block_type": "table", "content_text": "| x |", "page_number": 1, "sequence_index": 1, "metadata": {}},
    ]
    native = [
        {"block_type": "paragraph", "content_text": "the complete native paragraph", "page_number": 1, "sequence_index": 0, "metadata": {}},
    ]

    merged, quality = _apply_native_quality_gate(primary, native)

    assert quality["fallbackPages"] == [1]
    assert any(block["block_type"] == "table" for block in merged)
    fallback = next(block for block in merged if block["block_type"] == "paragraph")
    assert fallback["content_text"] == "the complete native paragraph"
    assert fallback["metadata"]["fallbackReason"] == "low_page_coverage"


def test_parser_routing_keeps_text_and_audio_local():
    assert uses_mineru(filename="notes.txt", content_type="text/plain") is False
    assert uses_mineru(filename="lecture.mp3", content_type="audio/mpeg") is False
    assert uses_mineru(filename="slides.pptx", content_type="application/octet-stream") is True
    assert uses_mineru(filename="book.xlsx", content_type="application/octet-stream") is True


def test_pptx_upload_validation_and_native_fallback():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "MinerU integration"
    slide.placeholders[1].text = "Native fallback content"
    buffer = BytesIO()
    presentation.save(buffer)
    data = buffer.getvalue()
    upload = UploadFile(BytesIO(data), filename="slides.pptx", headers=Headers())

    assert normalized_upload_content_type(upload).endswith("presentationml.presentation")
    assert validate_upload_file(upload) == len(data)
    blocks = parse_native_document_blocks(data, filename="slides.pptx")
    assert any(block["content_text"] == "MinerU integration" for block in blocks)
    assert all(block["metadata"]["ocrUsed"] is False for block in blocks)


def test_xlsx_upload_validation_and_native_fallback():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scores"
    sheet.append(["Name", "Score"])
    sheet.append(["Ada", 98])
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    upload = UploadFile(BytesIO(data), filename="scores.xlsx", headers=Headers())

    assert normalized_upload_content_type(upload).endswith("spreadsheetml.sheet")
    assert validate_upload_file(upload) == len(data)
    blocks = parse_native_document_blocks(data, filename="scores.xlsx")
    assert blocks[0]["block_type"] == "table"
    assert "Ada" in blocks[0]["content_text"]
    assert blocks[0]["structured_data"]["sheetName"] == "Scores"
