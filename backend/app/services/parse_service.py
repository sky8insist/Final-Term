from io import BytesIO
from hashlib import sha256
from pathlib import Path
import re

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from charset_normalizer import from_bytes

PDF_CONTENT_TYPE = "application/pdf"
TXT_CONTENT_TYPE = "text/plain"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
IMAGE_CONTENT_TYPES = {"image/png", "image/jpeg", "image/webp", "image/gif", "image/bmp", "image/jp2"}
AUDIO_CONTENT_TYPES = {"audio/mpeg", "audio/wav", "audio/x-wav", "audio/mp4", "audio/x-m4a"}


class DocumentParseError(ValueError):
    pass


PARSER_VERSION = "1.0.0"


def _block(*, block_type: str, text: str, sequence: int, parser: str,
           page_number: int | None = None, structured_data: dict | None = None,
           confidence: float = 1.0, metadata: dict | None = None,
           bounding_box: dict | None = None) -> dict:
    cleaned = _clean_text(text)
    identity = f"{parser}:{page_number}:{sequence}:{block_type}:{cleaned}".encode("utf-8")
    return {
        "block_type": block_type,
        "content_text": cleaned,
        "structured_data": structured_data or {},
        "page_number": page_number,
        "bounding_box": bounding_box,
        "start_time": None,
        "end_time": None,
        "sequence_index": sequence,
        "parser_name": parser,
        "parser_version": PARSER_VERSION,
        "confidence": confidence,
        "source_hash": sha256(identity).hexdigest(),
        "metadata": metadata or {},
    }


def _clean_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line).strip()


def _vision_content(result: dict, kind: str, fallback: str) -> str:
    parts: list[str] = []
    if kind == "table":
        table = result.get("table") or {}
        headers = [str(value) for value in table.get("headers", [])]
        rows = [[str(value) for value in row] for row in table.get("rows", [])]
        if headers:
            parts.append("| " + " | ".join(headers) + " |")
            parts.append("| " + " | ".join(["---"] * len(headers)) + " |")
            parts.extend("| " + " | ".join(row) + " |" for row in rows)
    parts.extend(filter(None, [str(result.get("ocrText", "")).strip(), str(result.get("summary", "")).strip()]))
    return "\n".join(parts) or fallback


def _looks_like_heading(text: str) -> bool:
    compact = text.strip()
    return len(compact) <= 100 and bool(re.match(
        r"^(第[一二三四五六七八九十百\d]+[章节篇部]|\d+(?:\.\d+){0,3}\s+|目录$|contents$)",
        compact, re.IGNORECASE,
    ))


def _parse_txt(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # GB18030 is the superset used for GBK/GB2312 course materials and is
        # more deterministic for short Chinese snippets than statistical
        # detectors, which can confuse them with Korean encodings.
        try:
            return data.decode("gb18030")
        except UnicodeDecodeError:
            pass
        match = from_bytes(data).best()
        if match is None or match.encoding is None or match.chaos > 0.25:
            raise DocumentParseError("TXT encoding could not be detected reliably")
        try:
            return str(match)
        except (UnicodeError, LookupError) as exc:
            raise DocumentParseError("TXT encoding could not be decoded") from exc


def _parse_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise DocumentParseError("PDF text could not be extracted") from exc

    return "\n".join(pages)


def _parse_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise DocumentParseError("DOCX text could not be extracted") from exc

    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def parse_document_bytes(
    data: bytes,
    *,
    filename: str | None = None,
    content_type: str | None = None,
) -> str:
    suffix = Path(filename or "").suffix.lower()

    if content_type == TXT_CONTENT_TYPE or suffix == ".txt":
        text = _parse_txt(data)
    elif content_type == PDF_CONTENT_TYPE or suffix == ".pdf":
        text = _parse_pdf(data)
    elif content_type == DOCX_CONTENT_TYPE or suffix == ".docx":
        text = _parse_docx(data)
    else:
        raise DocumentParseError("Unsupported file type")

    cleaned = _clean_text(text)
    if not cleaned:
        raise DocumentParseError("No extractable text found in this file")

    return cleaned


def parse_document_blocks(
    data: bytes, *, filename: str | None = None, content_type: str | None = None,
    audio_time_offset: float = 0.0, analyze_audio_transcript: bool = True,
    audio_segment_index: int | None = None,
) -> list[dict]:
    suffix = Path(filename or "").suffix.lower()
    blocks: list[dict] = []
    if content_type == PDF_CONTENT_TYPE or suffix == ".pdf":
        try:
            import fitz
            from collections import Counter
            from app.services.multimodal_service import analyze_image

            pdf = fitz.open(stream=data, filetype="pdf")
            if pdf.needs_pass:
                raise DocumentParseError("Encrypted PDF requires a password")
            layouts: list[tuple[object, list[tuple]]] = []
            margin_texts: Counter[str] = Counter()
            for page in pdf:
                page_blocks = [item for item in page.get_text("blocks") if _clean_text(str(item[4]))]
                layouts.append((page, page_blocks))
                height = max(float(page.rect.height), 1)
                for item in page_blocks:
                    text = _clean_text(str(item[4]))
                    if float(item[1]) <= height * 0.1 or float(item[3]) >= height * 0.9:
                        if len(text) <= 200:
                            margin_texts[text] += 1
            repeat_threshold = max(2, (len(layouts) + 1) // 2)
            repeated_margins = {text for text, count in margin_texts.items() if count >= repeat_threshold}

            for page_index, (page, page_blocks) in enumerate(layouts, start=1):
                width, height = max(float(page.rect.width), 1), max(float(page.rect.height), 1)
                kept_text = []
                for item in page_blocks:
                    text = _clean_text(str(item[4]))
                    if text in repeated_margins:
                        continue
                    kept_text.append(text)
                    blocks.append(_block(
                        block_type="heading" if _looks_like_heading(text) else "paragraph",
                        text=text, sequence=len(blocks),
                        parser="pymupdf", page_number=page_index,
                        bounding_box={
                            "x0": float(item[0]), "y0": float(item[1]),
                            "x1": float(item[2]), "y1": float(item[3]),
                            "pageWidth": width, "pageHeight": height,
                        },
                        metadata={"pageKind": "text", "repeatedMarginsRemoved": bool(repeated_margins)},
                    ))
                has_images = bool(page.get_images(full=True))
                needs_vision = not kept_text or has_images
                if needs_vision:
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    result = analyze_image(
                        pixmap.tobytes("png"), content_type="image/png",
                        filename=f"{filename or 'document.pdf'}#page={page_index}",
                    )
                    kind = result.get("kind", "image")
                    detail = result.get(kind) if isinstance(result.get(kind), dict) else {}
                    text = _vision_content(result, kind, f"第 {page_index} 页 OCR 未提取到可靠内容。")
                    blocks.append(_block(
                        block_type=kind, text=text, sequence=len(blocks), parser="pymupdf+vision-api",
                        page_number=page_index,
                        bounding_box={"x0": 0, "y0": 0, "x1": width, "y1": height,
                                      "pageWidth": width, "pageHeight": height},
                        structured_data={"ocrText": result.get("ocrText", ""), "summary": result.get("summary", ""), **detail},
                        confidence=result.get("confidence", 0.0),
                        metadata={"pageKind": "mixed" if kept_text else "scanned",
                                  "needsReview": result.get("confidence", 0.0) < 0.6},
                    ))
            pdf.close()
            blocks.sort(key=lambda block: (block.get("page_number") or 0, block["sequence_index"]))
            for sequence, block in enumerate(blocks):
                block["sequence_index"] = sequence
        except Exception as exc:
            raise DocumentParseError("PDF text could not be extracted") from exc
    elif content_type == DOCX_CONTENT_TYPE or suffix == ".docx":
        try:
            document = Document(BytesIO(data))
            table_index = 0
            image_index = 0
            for child in document.element.body.iterchildren():
                if isinstance(child, CT_P):
                    paragraph = Paragraph(child, document)
                    text = _clean_text(paragraph.text)
                    if text:
                        style = (paragraph.style.name if paragraph.style else "").lower()
                        block_type = "heading" if style.startswith("heading") or style.startswith("标题") else "paragraph"
                        blocks.append(_block(
                            block_type=block_type, text=text, sequence=len(blocks), parser="python-docx",
                            metadata={"style": paragraph.style.name if paragraph.style else None},
                        ))
                    for relationship_id in child.xpath(".//a:blip/@r:embed"):
                        relationship = document.part.rels.get(relationship_id)
                        if not relationship or not hasattr(relationship.target_part, "blob"):
                            continue
                        from app.services.multimodal_service import analyze_image
                        image_index += 1
                        image_result = analyze_image(
                            relationship.target_part.blob,
                            content_type=getattr(relationship.target_part, "content_type", "image/png"),
                            filename=f"{filename or 'document.docx'}#image={image_index}",
                        )
                        kind = image_result.get("kind", "image")
                        detail = image_result.get(kind) if isinstance(image_result.get(kind), dict) else {}
                        image_text = _vision_content(
                            image_result, kind, f"文档图片 {image_index} 未提取到可靠内容。",
                        )
                        blocks.append(_block(
                            block_type=kind, text=image_text, sequence=len(blocks), parser="python-docx+vision-api",
                            structured_data={"ocrText": image_result.get("ocrText", ""), **detail},
                            confidence=image_result.get("confidence", 0), metadata={"imageIndex": image_index},
                        ))
                elif isinstance(child, CT_Tbl):
                    table = Table(child, document)
                    rows = [[_clean_text(cell.text) for cell in row.cells] for row in table.rows]
                    if not rows:
                        continue
                    separator = ["---"] * len(rows[0])
                    markdown_rows = [rows[0], separator, *rows[1:]]
                    markdown = "\n".join("| " + " | ".join(row) + " |" for row in markdown_rows)
                    merged_cells = []
                    for row_index, row in enumerate(table.rows):
                        seen_xml_cells: set[int] = set()
                        for column_index, cell in enumerate(row.cells):
                            xml_identity = id(cell._tc)
                            if xml_identity in seen_xml_cells:
                                continue
                            seen_xml_cells.add(xml_identity)
                            properties = cell._tc.tcPr
                            grid_span = properties.gridSpan.val if properties is not None and properties.gridSpan is not None else 1
                            vertical = properties.vMerge.val if properties is not None and properties.vMerge is not None else None
                            if int(grid_span) > 1 or properties is not None and properties.vMerge is not None:
                                merged_cells.append({
                                    "row": row_index, "column": column_index,
                                    "columnSpan": int(grid_span),
                                    "verticalMerge": str(vertical) if vertical else "continue",
                                })
                    blocks.append(_block(
                        block_type="table", text=markdown, sequence=len(blocks), parser="python-docx",
                        structured_data={"rows": rows, "headers": rows[0], "tableIndex": table_index,
                                         "mergedCells": merged_cells},
                    ))
                    table_index += 1
        except Exception as exc:
            raise DocumentParseError("DOCX content could not be extracted") from exc
    elif content_type == TXT_CONTENT_TYPE or suffix == ".txt":
        text = _clean_text(_parse_txt(data))
        if text:
            blocks.append(_block(block_type="paragraph", text=text, sequence=0, parser="text"))
    elif content_type in IMAGE_CONTENT_TYPES or suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        from app.services.multimodal_service import analyze_image
        result = analyze_image(data, content_type=content_type or "image/jpeg", filename=filename or "image")
        kind = result.get("kind", "image")
        detail = result.get(kind) if isinstance(result.get(kind), dict) else {}
        text = _vision_content(result, kind, "图像内容未能可靠识别。")
        blocks.append(_block(
            block_type=kind, text=text, sequence=0, parser="vision-api",
            structured_data={"ocrText": result.get("ocrText", ""), "summary": result.get("summary", ""), **detail},
            confidence=result.get("confidence", 0.0), metadata={"needsReview": result.get("confidence", 0.0) < 0.6},
        ))
    elif content_type in AUDIO_CONTENT_TYPES or suffix in {".mp3", ".wav", ".m4a"}:
        from app.services.multimodal_service import analyze_transcript, transcribe_audio
        if content_type in {"audio/wav", "audio/x-wav"} or suffix == ".wav":
            try:
                import wave
                with wave.open(BytesIO(data), "rb") as audio_file:
                    local_duration = audio_file.getnframes() / max(audio_file.getframerate(), 1)
                from app.config.settings import settings
                if local_duration > settings.max_audio_minutes * 60:
                    raise DocumentParseError(f"Audio must be {settings.max_audio_minutes} minutes or shorter")
            except DocumentParseError:
                raise
            except (wave.Error, EOFError) as exc:
                raise DocumentParseError("WAV audio is invalid or empty") from exc
        result = transcribe_audio(data, content_type=content_type or "audio/mpeg", filename=filename or "audio")
        duration = result.get("duration")
        if duration is not None:
            from app.config.settings import settings
            if float(duration) > settings.max_audio_minutes * 60:
                raise DocumentParseError(f"Audio must be {settings.max_audio_minutes} minutes or shorter")
        blocks.extend(build_audio_blocks(
            result, audio_time_offset=audio_time_offset,
            analyze_audio_transcript=analyze_audio_transcript,
            audio_segment_index=audio_segment_index,
        ))
    else:
        raise DocumentParseError("Unsupported file type")
    if not blocks or not any(block["content_text"].strip() for block in blocks):
        raise DocumentParseError("No extractable text found in this file")
    return blocks


def parse_native_document_blocks(
    data: bytes, *, filename: str | None = None, content_type: str | None = None,
) -> list[dict]:
    """Extract native text without OCR or any external model.

    This is used only as MinerU's completeness gate/fallback. It deliberately
    ignores embedded images so PDF and Office requests cannot silently invoke
    OCR during the primary path.
    """
    suffix = Path(filename or "").suffix.lower()
    blocks: list[dict] = []
    if content_type == PDF_CONTENT_TYPE or suffix == ".pdf":
        try:
            import fitz
            pdf = fitz.open(stream=data, filetype="pdf")
            if pdf.needs_pass:
                raise DocumentParseError("Encrypted PDF requires a password")
            for page_number, page in enumerate(pdf, start=1):
                width, height = max(float(page.rect.width), 1), max(float(page.rect.height), 1)
                for item in page.get_text("blocks"):
                    text = _clean_text(str(item[4]))
                    if not text:
                        continue
                    blocks.append(_block(
                        block_type="heading" if _looks_like_heading(text) else "paragraph",
                        text=text, sequence=len(blocks), parser="pymupdf-native-fallback",
                        page_number=page_number,
                        bounding_box={
                            "x0": float(item[0]), "y0": float(item[1]),
                            "x1": float(item[2]), "y1": float(item[3]),
                            "pageWidth": width, "pageHeight": height,
                        },
                        metadata={"fallbackParser": True, "ocrUsed": False},
                    ))
            pdf.close()
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError("PDF native text could not be extracted") from exc
    elif content_type == DOCX_CONTENT_TYPE or suffix == ".docx":
        try:
            document = Document(BytesIO(data))
            for paragraph in document.paragraphs:
                text = _clean_text(paragraph.text)
                if text:
                    style = (paragraph.style.name if paragraph.style else "").lower()
                    blocks.append(_block(
                        block_type="heading" if style.startswith(("heading", "标题")) else "paragraph",
                        text=text, sequence=len(blocks), parser="python-docx-native-fallback",
                        metadata={"fallbackParser": True, "ocrUsed": False},
                    ))
            for table_index, table in enumerate(document.tables):
                rows = [[_clean_text(cell.text) for cell in row.cells] for row in table.rows]
                if rows:
                    markdown = "\n".join(
                        "| " + " | ".join(row) + " |"
                        for row in [rows[0], ["---"] * len(rows[0]), *rows[1:]]
                    )
                    blocks.append(_block(
                        block_type="table", text=markdown, sequence=len(blocks),
                        parser="python-docx-native-fallback",
                        structured_data={"rows": rows, "tableIndex": table_index},
                        metadata={"fallbackParser": True, "ocrUsed": False},
                    ))
        except Exception as exc:
            raise DocumentParseError("DOCX native text could not be extracted") from exc
    elif content_type == PPTX_CONTENT_TYPE or suffix == ".pptx":
        try:
            from pptx import Presentation
            presentation = Presentation(BytesIO(data))
            for slide_number, slide in enumerate(presentation.slides, start=1):
                for shape in slide.shapes:
                    text = _clean_text(getattr(shape, "text", ""))
                    if not text:
                        continue
                    is_title = shape == getattr(slide.shapes, "title", None)
                    blocks.append(_block(
                        block_type="heading" if is_title else "paragraph",
                        text=text, sequence=len(blocks), parser="python-pptx-native-fallback",
                        page_number=slide_number,
                        metadata={"fallbackParser": True, "ocrUsed": False, "slideNumber": slide_number},
                    ))
        except Exception as exc:
            raise DocumentParseError("PPTX native text could not be extracted") from exc
    elif content_type == XLSX_CONTENT_TYPE or suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
            for sheet_index, sheet in enumerate(workbook.worksheets):
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(value.strip() for value in values):
                        rows.append(values)
                if not rows:
                    continue
                width = max(len(row) for row in rows)
                normalized = [row + [""] * (width - len(row)) for row in rows]
                markdown = "\n".join(
                    "| " + " | ".join(row) + " |"
                    for row in [normalized[0], ["---"] * width, *normalized[1:]]
                )
                blocks.append(_block(
                    block_type="table", text=f"## {sheet.title}\n{markdown}",
                    sequence=len(blocks), parser="openpyxl-native-fallback",
                    page_number=sheet_index + 1,
                    structured_data={"sheetName": sheet.title, "rows": normalized},
                    metadata={"fallbackParser": True, "ocrUsed": False, "sheetIndex": sheet_index},
                ))
            workbook.close()
        except Exception as exc:
            raise DocumentParseError("XLSX native data could not be extracted") from exc
    else:
        raise DocumentParseError("No native fallback parser is available for this file type")
    return blocks


def build_audio_blocks(
    result: dict, *, audio_time_offset: float = 0.0,
    analyze_audio_transcript: bool = True, audio_segment_index: int | None = None,
) -> list[dict]:
    from app.services.multimodal_service import analyze_transcript

    blocks: list[dict] = []
    segments = result.get("segments") or []
    full_transcript = str(
        result.get("text") or " ".join(str(item.get("text", "")) for item in segments)
    ).strip()
    analysis = analyze_transcript(full_transcript) if full_transcript and analyze_audio_transcript else {}
    if segments:
        for segment in segments:
            text = str(segment.get("text", "")).strip()
            if not text:
                continue
            block = _block(block_type="audio", text=text, sequence=len(blocks), parser="transcription-api")
            local_start = max(float(segment.get("start", 0)), 0.0)
            local_end = max(float(segment.get("end", local_start)), local_start)
            block["start_time"] = audio_time_offset + local_start
            block["end_time"] = audio_time_offset + local_end
            block["structured_data"] = {
                "language": result.get("language"), "speaker": segment.get("speaker"),
                "audioSegmentIndex": audio_segment_index,
            }
            if not blocks:
                block["structured_data"]["analysis"] = analysis
            blocks.append(block)
    elif result.get("text"):
        block = _block(block_type="audio", text=result["text"], sequence=0, parser="transcription-api")
        block["start_time"] = audio_time_offset
        block["end_time"] = (
            audio_time_offset + float(result["duration"])
            if result.get("duration") is not None else None
        )
        block["structured_data"] = {
            "language": result.get("language"), "analysis": analysis,
            "audioSegmentIndex": audio_segment_index,
        }
        blocks.append(block)
    return blocks


def parse_document(path: str | Path) -> str:
    file_path = Path(path)
    return parse_document_bytes(file_path.read_bytes(), filename=file_path.name)
