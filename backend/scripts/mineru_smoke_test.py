"""Live MinerU format smoke test using small synthetic, non-sensitive files."""

from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation

from app.services.mineru_parser import parse_material


def _docx() -> bytes:
    document = Document()
    document.add_heading("MinerU Word smoke test", level=1)
    document.add_paragraph("Document parsing should preserve this sentence.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "MinerU PowerPoint smoke test"
    slide.placeholders[1].text = "Slide parsing should preserve this sentence."
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scores"
    sheet.append(["Student", "Score"])
    sheet.append(["Ada", 98])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _png() -> bytes:
    image = Image.new("RGB", (900, 240), "white")
    ImageDraw.Draw(image).text((40, 90), "MinerU image smoke test 2026", fill="black")
    stream = BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def main() -> None:
    cases = [
        ("smoke.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", _docx()),
        ("smoke.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", _pptx()),
        ("smoke.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _xlsx()),
        ("smoke.png", "image/png", _png()),
    ]
    report = []
    for index, (filename, content_type, data) in enumerate(cases):
        result = parse_material(
            data, filename=filename, content_type=content_type,
            material_id=f"synthetic-live-smoke-{index}",
        )
        metadata = result.metadata or {}
        report.append({
            "file": filename,
            "provider": result.provider,
            "blocks": len(result.blocks),
            "ocrUsed": metadata.get("ocrUsed"),
            "fallbackReason": metadata.get("fallbackReason"),
            "qualityGate": metadata.get("qualityGate"),
        })
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
